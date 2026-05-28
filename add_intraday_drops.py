# -*- coding: utf-8 -*-
"""把今年以来 UPRO 当日最低价跌破前收≥3% 的交易日，作为新工作表追加到交易复盘文档。"""
import sys
import akshare as ak
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

sys.stdout.reconfigure(encoding="utf-8")

F = r"C:\project\股票\标普3倍做多日期20150101-20260527交易复盘_已匹配.xlsx"
SHEET = "盘中跌破3%(2026)"

# 1. 取数 + 筛选
df = ak.stock_us_daily(symbol="UPRO", adjust="")
df["date"] = pd.to_datetime(df["date"])
df = df[df["date"] >= "2026-01-01"].sort_values("date").reset_index(drop=True)
df["prev_close"] = df["close"].shift(1)
df = df.dropna().reset_index(drop=True)
df["low_pct"] = (df["low"] - df["prev_close"]) / df["prev_close"] * 100
df["close_pct"] = (df["close"] - df["prev_close"]) / df["prev_close"] * 100
sel = df[df["low_pct"] <= -3.0].reset_index(drop=True)

# 2. 写入新工作表（已存在则覆盖重建）
wb = openpyxl.load_workbook(F)
if SHEET in wb.sheetnames:
    del wb[SHEET]
ws = wb.create_sheet(SHEET)

ws["A1"] = "今年以来 UPRO 当日最低价跌破前收 ≥3% 的交易日（数据源 akshare，截至 2026-05-28）"
ws["A1"].font = Font(bold=True, size=12)
ws.merge_cells("A1:G1")

headers = ["日期", "前收", "当日最低", "最低跌幅", "收盘", "收盘跌幅", "类型"]
HR = 2
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=HR, column=c, value=h)
    cell.font = Font(color="FFFFFF", bold=True)
    cell.fill = PatternFill("solid", fgColor="305496")
    cell.alignment = Alignment(horizontal="center")

red = Font(color="C00000")
green = Font(color="008000")
wick_fill = PatternFill("solid", fgColor="FFF2CC")  # 盘中插针行高亮

for i, r in sel.iterrows():
    row = HR + 1 + i
    is_true = r["close_pct"] <= -3.0
    ws.cell(row=row, column=1, value=r["date"].strftime("%Y-%m-%d"))
    ws.cell(row=row, column=2, value=round(r["prev_close"], 2))
    ws.cell(row=row, column=3, value=round(r["low"], 2))
    c4 = ws.cell(row=row, column=4, value=round(r["low_pct"], 2)); c4.number_format = '0.00"%"'; c4.font = red
    ws.cell(row=row, column=5, value=round(r["close"], 2))
    c6 = ws.cell(row=row, column=6, value=round(r["close_pct"], 2)); c6.number_format = '0.00"%"'
    c6.font = red if r["close_pct"] < 0 else green
    ws.cell(row=row, column=7, value=("真跌" if is_true else "盘中插针"))
    if not is_true:
        for cc in range(1, 8):
            ws.cell(row=row, column=cc).fill = wick_fill

n_true = int((sel["close_pct"] <= -3.0).sum())
SR = HR + 1 + len(sel) + 1
ws.cell(row=SR, column=1,
        value=f"命中 {len(sel)} 天：真跌(收盘也跌破3%) {n_true} 天，盘中插针后收回 {len(sel)-n_true} 天").font = Font(italic=True)

for i, w in enumerate([12, 9, 10, 10, 9, 10, 12], 1):
    ws.column_dimensions[chr(64 + i)].width = w

wb.save(F)
print(f"已写入工作表 '{SHEET}'：{len(sel)} 行（真跌 {n_true}，插针 {len(sel)-n_true}）")
print("当前工作表:", wb.sheetnames)
