"""LIFO 匹配买入与卖出，把卖出价/涨幅回填到每笔买入行。"""
import xlrd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from collections import deque
import sys

sys.stdout.reconfigure(encoding='utf-8')

SRC = r"C:\project\股票\标普3倍做多日期20150101-20260527交易复盘.xls"
DST = r"C:\project\股票\标普3倍做多日期20150101-20260527交易复盘_已匹配.xlsx"

wb_src = xlrd.open_workbook(SRC)
sh = wb_src.sheet_by_name('原始交易记录')

header = [sh.cell_value(0, c) for c in range(sh.ncols)]
rows = []
for r in range(1, sh.nrows):
    rows.append([sh.cell_value(r, c) for c in range(sh.ncols)])

COL_QTY, COL_PRICE, COL_SIDE = 4, 5, 3
COL_SELL_AMT, COL_SELL_PCT = 11, 12

# 每笔买入累计的卖出股数 / 加权卖出金额
buy_match = {}  # idx -> {"sold_qty": float, "sold_value": float}
stack = deque()  # 每个元素: [row_idx, remaining_qty, buy_price]
warnings = []

for i, row in enumerate(rows):
    side = str(row[COL_SIDE]).strip()
    qty = float(row[COL_QTY])
    price = float(row[COL_PRICE])
    if side == '买入':
        stack.append([i, qty, price])
        buy_match[i] = {"sold_qty": 0.0, "sold_value": 0.0, "buy_price": price, "buy_qty": qty}
    elif side == '卖出':
        remaining = qty
        while remaining > 1e-9:
            if not stack:
                warnings.append(f"row {i+2}: 卖出 {remaining} 股无买入可匹配（持仓不足）")
                break
            top = stack[-1]
            take = min(top[1], remaining)
            buy_match[top[0]]["sold_qty"] += take
            buy_match[top[0]]["sold_value"] += take * price
            top[1] -= take
            remaining -= take
            if top[1] <= 1e-9:
                stack.pop()
    else:
        warnings.append(f"row {i+2}: 未知买卖方向: {side}")

# 写回
wb_out = Workbook()
ws = wb_out.active
ws.title = '原始交易记录'

# header
ws.append(header)
header_fill = PatternFill('solid', fgColor='305496')
header_font = Font(color='FFFFFF', bold=True)
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

buy_fill = PatternFill('solid', fgColor='E2EFDA')   # 浅绿 = 买入
sell_fill = PatternFill('solid', fgColor='FCE4D6')  # 浅橙 = 卖出
gain_font = Font(color='C00000', bold=True)
loss_font = Font(color='008000', bold=True)

for i, row in enumerate(rows):
    out = list(row)
    side = str(row[COL_SIDE]).strip()
    if side == '买入':
        m = buy_match[i]
        if m["sold_qty"] > 1e-9:
            avg_sell = m["sold_value"] / m["sold_qty"]
            pct = (avg_sell - m["buy_price"]) / m["buy_price"] * 100
            out[COL_SELL_AMT] = round(avg_sell, 4)
            out[COL_SELL_PCT] = round(pct, 4)
            if m["sold_qty"] < m["buy_qty"] - 1e-9:
                out[COL_SELL_AMT] = f"{round(avg_sell,4)} (部分:{m['sold_qty']:.0f}/{m['buy_qty']:.0f})"
        else:
            out[COL_SELL_AMT] = '持仓中'
            out[COL_SELL_PCT] = ''
    ws.append(out)

    excel_row = i + 2
    fill = buy_fill if side == '买入' else sell_fill
    for c_idx in range(1, len(out) + 1):
        ws.cell(row=excel_row, column=c_idx).fill = fill

    if side == '买入' and isinstance(out[COL_SELL_PCT], (int, float)):
        pct = out[COL_SELL_PCT]
        cell = ws.cell(row=excel_row, column=COL_SELL_PCT + 1)
        cell.number_format = '0.00"%"'
        cell.font = gain_font if pct >= 0 else loss_font
        amt_cell = ws.cell(row=excel_row, column=COL_SELL_AMT + 1)
        if isinstance(out[COL_SELL_AMT], (int, float)):
            amt_cell.number_format = '0.0000'

# 列宽
widths = [10, 6, 8, 6, 10, 12, 6, 12, 12, 14, 10, 14, 12, 12]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[chr(64 + i) if i <= 26 else 'A' + chr(64 + i - 26)].width = w

# 在末尾追加汇总
sum_row = len(rows) + 3
ws.cell(row=sum_row, column=1, value='【LIFO 匹配汇总】').font = Font(bold=True, size=12)
total_pnl = 0.0
total_cost = 0.0
matched = 0
for i, row in enumerate(rows):
    side = str(row[COL_SIDE]).strip()
    if side != '买入':
        continue
    m = buy_match[i]
    if m["sold_qty"] > 1e-9:
        pnl = m["sold_value"] - m["buy_price"] * m["sold_qty"]
        total_pnl += pnl
        total_cost += m["buy_price"] * m["sold_qty"]
        matched += 1

ws.cell(row=sum_row + 1, column=1, value=f'已匹配买入笔数：{matched}')
ws.cell(row=sum_row + 2, column=1, value=f'累计盈亏（USD）：{total_pnl:.2f}')
ws.cell(row=sum_row + 3, column=1, value=f'累计成本（USD）：{total_cost:.2f}')
if total_cost:
    ws.cell(row=sum_row + 4, column=1, value=f'整体收益率：{total_pnl/total_cost*100:.2f}%')
if stack:
    remain_qty = sum(t[1] for t in stack)
    ws.cell(row=sum_row + 5, column=1, value=f'未平仓股数：{remain_qty:.0f}')
for w in warnings:
    print('WARN:', w)

wb_out.save(DST)
print(f"已生成: {DST}")
print(f"已匹配买入笔数: {matched}, 累计盈亏 USD: {total_pnl:.2f}")
if stack:
    print(f"未平仓: {[(t[0]+2, t[1], t[2]) for t in stack]}")
